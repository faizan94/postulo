"""Pick the reader by file type: text PDF -> pypdf, .docx -> python-docx, text -> as is."""

import io
from dataclasses import dataclass, field

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.config import Settings
from app.errors import BadInputError
from app.services.intake.clean_text import clean

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 10
MAX_TEXT = 20_000
SCAN_CHARS_PER_PAGE = 100
SCAN_MESSAGE = (
    "This looks like a scan or photo. Please upload a text PDF or .docx, or paste the text."
)


@dataclass
class ExtractResult:
    text: str
    source_type: str
    label: str
    ocr_used: bool = False
    warnings: list[str] = field(default_factory=list)


def _pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise BadInputError("The PDF is password-protected.")
        pages = [p.extract_text() or "" for p in reader.pages[:MAX_PAGES]]
    except PdfReadError as e:
        raise BadInputError(f"Could not read the PDF: {e}") from e
    text = "\n\n".join(p.strip() for p in pages if p.strip())
    if len(text) < SCAN_CHARS_PER_PAGE * max(1, len(pages)):
        raise BadInputError(SCAN_MESSAGE)
    return text


def _docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:  # python-docx raises several types for broken files
        raise BadInputError("Could not read the Word file. Save it again as .docx or PDF.") from e
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def extract(
    settings: Settings,
    *,
    filename: str | None = None,
    data: bytes | None = None,
    text: str | None = None,
) -> ExtractResult:
    if text is not None and text.strip():
        raw, source, label = text, "text", "Pasted text"
    elif data is not None:
        if not data:
            raise BadInputError("The uploaded file is empty.")
        if len(data) > MAX_BYTES:
            raise BadInputError("The file is larger than 10 MB.")
        name = (filename or "cv").strip() or "cv"
        lower = name.lower()
        label = name
        if data[:5] == b"%PDF-":
            raw, source = _pdf(data), "pdf"
        elif data[:2] == b"PK" and lower.endswith(".docx"):
            raw, source = _docx(data), "docx"
        elif lower.endswith(".doc"):
            raise BadInputError("Old .doc files aren't supported. Save it as PDF or .docx.")
        elif lower.endswith((".png", ".jpg", ".jpeg", ".webp")):
            if not settings.ocr_enabled:
                raise BadInputError(SCAN_MESSAGE)
            raise BadInputError("Image reading is not available yet.")
        else:
            raise BadInputError("Upload a PDF or .docx file, or paste the text.")
    else:
        raise BadInputError("Upload a file or paste your CV text.")

    warnings: list[str] = []
    cleaned = clean(raw)
    if len(cleaned) > MAX_TEXT:
        cleaned = cleaned[:MAX_TEXT]
        warnings.append("Your CV was long, so only the first 20,000 characters are used.")
    if len(cleaned) < 50:
        raise BadInputError("We found almost no text in this CV.")
    return ExtractResult(text=cleaned, source_type=source, label=label, warnings=warnings)
